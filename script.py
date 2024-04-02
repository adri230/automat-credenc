#########################
#pip install python-docx#
#pip install openpyxl   #
#########################

from docx import Document # Se importa Document de la libreria docx para modificar el docx
import openpyxl # Se importa la libreria de openpyxl para leer excel
import os # Se importa la biblioteca os

#############
##Variables##
#############
doc = input('Nombre del fichero docx: ')
excel=input('Nombre del fichero xlsx: ')

doc=doc+".docx"
excel=excel+".xlsx"
try:
    doc= Document(doc)
    excel= openpyxl.load_workbook(excel)
except FileNotFoundError:
    print("Nombre de documento o excel no encontrado") #Si no se encuentran los archivos
except Exception as e:
    print("Error al abrir el archivo: ",e) #Error de otro tipo

contador=0 # Se empieza un contador
nombre= 'NOMBRE'+str(contador) 
dnip='DNI'+str(contador)
recint='Recinto'+str(contador)
date='Fecha'+str(contador)
excel1= excel.active # Usado para hacer que lea el excel
nombres=[] # Array de nombres dni recinto y fecha
dni=[]
recinto=[]
fecha=[]

##########################################
#For para ir cogiendo los datos del excel#
##########################################
try:
    for row in excel1.iter_rows(values_only=True):
        nombres.append(row[0])
        dni.append(row[1])
        recinto.append(row[2])
        fecha.append(row[3])
except Exception as e:
    print("Se ha producido un error al leer los datos en el archivo Excel: ",e)

###############################
#While para ir cambiando datos#    
###############################

while contador <= 5: # Contador con el número de tarjetas por hoja
    if contador >= len(nombres): #Se comprueba si se llega al final de la lista de nombres y se sale
        print("El excel ha llegado a su fin")
        break 

    for paragraph in doc.paragraphs: # Reemplazar las palabras en el texto del documento
        if nombre in paragraph.text:
            paragraph.text = paragraph.text.replace(nombre, nombres[contador])
        if dnip in paragraph.text:
            paragraph.text = paragraph.text.replace(dnip, dni[contador])
        if recint in paragraph.text:
            paragraph.text = paragraph.text.replace(recint,recinto[contador])
        if date in paragraph.text:
            paragraph.text = paragraph.text.replace(date,fecha[contador])

    contador+=1 # Se actualiza el contador
    nombre='NOMBRE'+str(contador)
    dnip='DNI'+str(contador)
    recint='Recinto'+str(contador)
    date='Fecha'+str(contador)


doc.save('final.docx') # Guardar los cambios en el archivo .docx
